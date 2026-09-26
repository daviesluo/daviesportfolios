#!/usr/bin/env python3
"""Touch-spread history from Bitget's public level-1 depth files (hist/depth/<SYM>/*.zip),
each an .xlsx of (timestamp s, ask, bid, ask qty, bid qty) snapshots for one day.

Per book and month: files (days), snapshots, the median of the daily median spread (bps),
the median of the daily 90th percentile, the share of snapshots wider than 10 / 30 bps,
and the median top-of-book size on the smaller side (in base units).
Days are the ones pulled (Wednesdays and Saturdays by default). Writes out/depth_summary.json.
"""
import io, json, os, sys, glob, zipfile, datetime, statistics as st
from collections import defaultdict
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))


def read_file(path):
    z = zipfile.ZipFile(path)
    snaps = []
    for n in z.namelist():
        b = z.read(n)
        if n.endswith(".xlsx"):
            wb = openpyxl.load_workbook(io.BytesIO(b), read_only=True)
            for row in wb[wb.sheetnames[0]].iter_rows(values_only=True):
                if row[0] == "timestamp" or row[0] is None:
                    continue
                try:
                    t, a, bb, aq, bq = int(row[0]), float(row[1]), float(row[2]), float(row[3]), float(row[4])
                except (TypeError, ValueError):
                    continue
                snaps.append((t, a, bb, aq, bq))
        else:
            for line in b.decode().splitlines()[1:]:
                p = line.split(",")
                try:
                    snaps.append((int(float(p[0])), float(p[1]), float(p[2]), float(p[3]), float(p[4])))
                except (ValueError, IndexError):
                    pass
    return sorted(snaps)


def pct(xs, q):
    xs = sorted(xs)
    return xs[min(len(xs) - 1, int(q * (len(xs) - 1) + 0.5))]


def main():
    syms = sys.argv[1:] or sorted(os.listdir(os.path.join(HERE, "hist", "depth")))
    out = {}
    for sym in syms:
        mon = defaultdict(lambda: {"days": 0, "snaps": 0, "dmed": [], "dp90": [], "w10": 0, "w30": 0, "qmin": [],
                                   "wd": [], "we": []})
        for f in sorted(glob.glob(os.path.join(HERE, "hist", "depth", sym, "*.zip"))):
            try:
                snaps = read_file(f)
            except Exception as e:
                print("skip", f, e)
                continue
            snaps = [s for s in snaps if s[1] > 0 and s[2] > 0 and s[1] >= s[2]]
            if not snaps:
                continue
            day = os.path.basename(f).split("_")[1]
            mk = day[:7]
            spreads = [(a - b) / ((a + b) / 2) * 1e4 for _, a, b, _, _ in snaps]
            d = mon[mk]
            d["days"] += 1
            d["snaps"] += len(spreads)
            d["dmed"].append(st.median(spreads))
            d["dp90"].append(pct(spreads, 0.9))
            d["w10"] += sum(x > 10 for x in spreads)
            d["w30"] += sum(x > 30 for x in spreads)
            d["qmin"].append(st.median(min(aq, bq) for _, _, _, aq, bq in snaps))
            wk = datetime.date.fromisoformat(day).weekday()
            (d["we"] if wk >= 5 else d["wd"]).append(st.median(spreads))
        rows = {}
        for mk in sorted(mon):
            d = mon[mk]
            rows[mk] = {"days": d["days"], "snapshots": d["snaps"],
                        "daily_median_spread_bps_median": round(st.median(d["dmed"]), 2),
                        "daily_p90_spread_bps_median": round(st.median(d["dp90"]), 2),
                        "share_snapshots_gt_10bps": round(d["w10"] / d["snaps"], 3),
                        "share_snapshots_gt_30bps": round(d["w30"] / d["snaps"], 3),
                        "weekday_days_median_bps": round(st.median(d["wd"]), 2) if d["wd"] else None,
                        "weekend_days_median_bps": round(st.median(d["we"]), 2) if d["we"] else None,
                        "top_qty_smaller_side_median_base": round(st.median(d["qmin"]), 2)}
        out[sym] = rows
        print(f"\n== {sym}")
        print(f"{'month':8} {'days':>4} {'snaps':>7} {'medSpr':>7} {'p90':>7} {'>10bp':>6} {'>30bp':>6} {'wkday':>6} {'wkend':>6} {'topQty':>9}")
        for mk, r in rows.items():
            g = lambda k: "-" if r[k] is None else f"{r[k]:.1f}"
            print(f"{mk:8} {r['days']:>4} {r['snapshots']:>7} {r['daily_median_spread_bps_median']:>7.2f} "
                  f"{r['daily_p90_spread_bps_median']:>7.2f} {r['share_snapshots_gt_10bps']:>6.3f} {r['share_snapshots_gt_30bps']:>6.3f} "
                  f"{g('weekday_days_median_bps'):>6} {g('weekend_days_median_bps'):>6} {r['top_qty_smaller_side_median_base']:>9.1f}")
    os.makedirs(os.path.join(HERE, "out"), exist_ok=True)
    path = os.path.join(HERE, "out", "depth_summary.json")
    old = json.load(open(path)) if os.path.exists(path) else {}
    old.update(out)
    json.dump(old, open(path, "w"), indent=1)


if __name__ == "__main__":
    main()
